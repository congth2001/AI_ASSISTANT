from datetime import date, datetime
from decimal import Decimal
from io import BytesIO

from openpyxl import load_workbook

from src.domain.entities.dashboard_analytics import (
    OrderLedgerExportData,
    OrderLedgerExportFilters,
    OrderLedgerExportItem,
)
from src.infrastructure.exporters.order_ledger_workbook_exporter import (
    OpenpyxlOrderLedgerWorkbookExporter,
)


def item(voucher, reason, running_debt, product_name="Sản phẩm A"):
    return OrderLedgerExportItem(
        customer_id="c1",
        customer_name="Khách A",
        customer_address="Số 12 đường Trung Tâm",
        occurred_at=datetime(2025, 3, 10, 9, 30),
        voucher_number=voucher,
        reason=reason,
        product_name=product_name if reason != "Thu nợ" else None,
        unit_name="Cái" if reason != "Thu nợ" else None,
        unit_price=Decimal("50000.50") if reason != "Thu nợ" else Decimal("0"),
        quantity=Decimal("2.500") if reason != "Thu nợ" else Decimal("0"),
        payment_amount=Decimal("20000.40") if reason != "Nhập trả" else Decimal("0"),
        debt_amount=Decimal("80500.10") if reason == "Xuất bán" else Decimal("-20000.40") if reason == "Nhập trả" else Decimal("0"),
        running_debt=Decimal(running_debt),
    )


def build(items):
    content = OpenpyxlOrderLedgerWorkbookExporter().build(
        OrderLedgerExportData(
            customer_id="c1",
            customer_name="Khách A",
            customer_address="Số 12 đường Trung Tâm",
            items=tuple(items),
        ),
        OrderLedgerExportFilters(
            date_from=date(2025, 3, 1),
            date_to=date(2025, 3, 31),
        ),
    )
    return load_workbook(BytesIO(content), data_only=False)


def test_workbook_matches_reference_layout_and_collapses_repeated_event_fields():
    workbook = build([
        item("XB-001", "Xuất bán", "80500.10", "Sản phẩm A"),
        item("XB-001", "Xuất bán", "80500.10", "Sản phẩm B"),
        item("TN-001", "Thu nợ", "60500.10"),
    ])

    sheet = workbook["Khách A_DoiTac"]
    assert sheet["A1"].value == "Tên đối tác:"
    assert sheet["B1"].value == "Khách A"
    assert sheet["A2"].value == "Địa chỉ:"
    assert sheet["B2"].value == "Số 12 đường Trung Tâm"
    assert sheet["A3"].value == "Ngày xuất file:"
    assert isinstance(sheet["B3"].value, datetime)
    assert sheet["A4"].value is None
    assert [sheet.cell(row=6, column=column).value for column in range(1, 11)] == [
        "Ngày", "Số phiếu", "Lý do", "Tên hàng", "ĐVT", "Đơn giá",
        "Số lượng", "Thanh toán", "Ghi nợ", "Nợ cộng dồn",
    ]
    assert sheet["B7"].value == "XB-001"
    assert sheet["I7"].value == 80500
    assert sheet["J7"].value == 80500
    assert sheet["A8"].value is None
    assert sheet["B8"].value is None
    assert sheet["C8"].value is None
    assert sheet["D8"].value == "Sản phẩm B"
    assert sheet["H8"].value is None
    assert sheet["I8"].value is None
    assert sheet["J8"].value is None
    assert sheet["C9"].value == "Thu nợ"
    assert sheet["A7"].number_format == "d/m/yyyy h:mm AM/PM"
    assert sheet["F7"].number_format == OpenpyxlOrderLedgerWorkbookExporter.ACCOUNTING_FORMAT
    assert sheet.freeze_panes is None
    assert sheet.auto_filter.ref is None


def test_empty_workbook_keeps_customer_metadata_and_is_readable():
    workbook = build([])
    sheet = workbook["Khách A_DoiTac"]

    assert sheet["B1"].value == "Khách A"
    assert sheet["A7"].value == "Không có giao dịch phù hợp với bộ lọc."
