from datetime import date
from decimal import Decimal
from io import BytesIO

from openpyxl import load_workbook

from src.domain.entities.dashboard_analytics import (
    CustomerDebtExportFilters,
    CustomerDebtExportItem,
    CustomerDebtSheetMode,
)
from src.infrastructure.exporters.customer_debt_workbook_exporter import (
    OpenpyxlCustomerDebtWorkbookExporter,
)


def item(name, ward, village, debt):
    return CustomerDebtExportItem(
        key=name,
        label=name,
        ward_name=ward,
        village_name=village,
        address_detail=f"Địa chỉ {name}",
        current_debt=Decimal(debt),
    )


def build_workbook(items, filters):
    content = OpenpyxlCustomerDebtWorkbookExporter().build(items, filters)
    return load_workbook(BytesIO(content), data_only=False)


def test_single_sheet_is_sorted_by_debt_and_uses_numeric_money_cells():
    workbook = build_workbook(
        [
            item("Khách B", "Xã 1", "Thôn B", "20.00"),
            item("Khách A", "Xã 1", "Thôn A", "100.50"),
            item("Khách C", "Xã 2", "Thôn C", "-10.00"),
        ],
        CustomerDebtExportFilters(
            date_from=date(2025, 3, 1),
            date_to=date(2025, 3, 31),
        ),
    )

    assert workbook.sheetnames == ["All"]
    sheet = workbook.active
    assert sheet.freeze_panes == "A5"
    assert sheet["B5"].value == "Khách A"
    assert sheet["B6"].value == "Khách B"
    assert sheet["B7"].value == "Khách C"
    assert sheet["F5"].value == 101
    assert sheet["F7"].value == -10
    assert sheet["F5"].number_format == '#,##0;[Red]-#,##0'
    assert sheet["F8"].value == "=SUM(F5:F7)"
    assert sheet.auto_filter.ref == "A4:F7"
    assert "01/03/2025 - 31/03/2025" in sheet["A2"].value


def test_village_mode_creates_safe_unique_sheets_sorted_by_debt():
    workbook = build_workbook(
        [
            item("Khách thấp", "Xã A", "Thôn/Một", "10"),
            item("Khách cao", "Xã A", "Thôn/Một", "50"),
            item("Khách khác", "Xã B", "Thôn/Một", "20"),
        ],
        CustomerDebtExportFilters(
            ward_names=("Xã A", "Xã B"),
            sheet_mode=CustomerDebtSheetMode.VILLAGE,
        ),
    )

    assert len(workbook.sheetnames) == 3
    assert workbook.sheetnames[0] == "All"
    assert all("/" not in name and len(name) <= 31 for name in workbook.sheetnames)
    assert [workbook["All"][f"B{row}"].value for row in range(5, 8)] == [
        "Khách cao",
        "Khách khác",
        "Khách thấp",
    ]
    first_sheet = workbook[workbook.sheetnames[1]]
    assert [first_sheet["B5"].value, first_sheet["B6"].value] == [
        "Khách cao",
        "Khách thấp",
    ]


def test_ward_mode_groups_rows_by_village_then_sorts_debt_within_village():
    workbook = build_workbook(
        [
            item("B thấp", "Xã A", "Thôn B", "10"),
            item("A thấp", "Xã A", "Thôn A", "20"),
            item("A cao", "Xã A", "Thôn A", "90"),
        ],
        CustomerDebtExportFilters(sheet_mode=CustomerDebtSheetMode.WARD),
    )

    assert workbook.sheetnames[0] == "All"
    sheet = workbook["Xã A"]
    assert [sheet[f"B{row}"].value for row in range(5, 8)] == [
        "A cao",
        "A thấp",
        "B thấp",
    ]
    assert [sheet[f"D{row}"].value for row in range(5, 8)] == [
        "Thôn A",
        "Thôn A",
        "Thôn B",
    ]


def test_empty_export_still_contains_a_readable_sheet():
    workbook = build_workbook([], CustomerDebtExportFilters())

    assert workbook.sheetnames == ["All"]
    assert workbook.active["A5"].value == "Tổng cộng (0 khách hàng)"
    assert workbook.active["F5"].value == 0


def test_export_escapes_formula_like_customer_text():
    workbook = build_workbook(
        [item("=HYPERLINK(\"bad\")", "Xã A", "Thôn A", "10")],
        CustomerDebtExportFilters(),
    )

    assert workbook.active["B5"].data_type == "s"
    assert workbook.active["B5"].value == "'=HYPERLINK(\"bad\")"
