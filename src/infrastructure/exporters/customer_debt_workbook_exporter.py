"""OpenPyXL adapter for customer debt workbooks."""

from collections import defaultdict
from datetime import date
from decimal import Decimal, ROUND_HALF_UP
from io import BytesIO
import re

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from src.domain.entities.dashboard_analytics import (
    CustomerDebtExportFilters,
    CustomerDebtExportItem,
    CustomerDebtSheetMode,
)
from src.domain.services.i_customer_debt_workbook_exporter import (
    ICustomerDebtWorkbookExporter,
)


class OpenpyxlCustomerDebtWorkbookExporter(ICustomerDebtWorkbookExporter):
    MAX_DATA_ROWS = 1_048_570
    HEADERS = (
        "STT",
        "Khách hàng",
        "Phường/xã",
        "Thôn/xóm",
        "Địa chỉ",
        "Công nợ",
    )

    def build(
        self,
        items: list[CustomerDebtExportItem],
        filters: CustomerDebtExportFilters,
    ) -> bytes:
        workbook = Workbook()
        workbook.remove(workbook.active)
        used_names: set[str] = set()

        for raw_name, sheet_items in self._group_items(items, filters):
            if len(sheet_items) > self.MAX_DATA_ROWS:
                raise ValueError(f"Sheet '{raw_name}' exceeds the Excel row limit")
            sheet_name = self._unique_sheet_name(raw_name, used_names)
            worksheet = workbook.create_sheet(sheet_name)
            self._populate_sheet(worksheet, sheet_items, filters)

        workbook.calculation.fullCalcOnLoad = True
        workbook.calculation.forceFullCalc = True
        output = BytesIO()
        workbook.save(output)
        return output.getvalue()

    def _group_items(
        self,
        items: list[CustomerDebtExportItem],
        filters: CustomerDebtExportFilters,
    ) -> list[tuple[str, list[CustomerDebtExportItem]]]:
        all_items = self._sort_by_debt(items)
        if filters.sheet_mode == CustomerDebtSheetMode.SINGLE:
            return [("All", all_items)]

        groups: dict[tuple[str, ...], list[CustomerDebtExportItem]] = defaultdict(list)
        if filters.sheet_mode == CustomerDebtSheetMode.VILLAGE:
            for item in items:
                groups[(item.ward_name or "Không rõ xã", item.village_name or "Không rõ thôn")].append(item)
            return [("All", all_items), *[
                (f"{village} - {ward}", self._sort_by_debt(group))
                for (ward, village), group in sorted(groups.items())
            ]]

        for item in items:
            groups[(item.ward_name or "Không rõ xã",)].append(item)
        return [("All", all_items), *[
            (ward, sorted(group, key=lambda item: (
                (item.village_name or "").casefold(),
                -item.current_debt,
                item.label.casefold(),
            )))
            for (ward,), group in sorted(groups.items())
        ]]

    @staticmethod
    def _sort_by_debt(
        items: list[CustomerDebtExportItem],
    ) -> list[CustomerDebtExportItem]:
        return sorted(items, key=lambda item: (-item.current_debt, item.label.casefold()))

    @staticmethod
    def _unique_sheet_name(raw_name: str, used_names: set[str]) -> str:
        cleaned = re.sub(r"[\\/*?:\[\]]", "-", raw_name).strip(" '") or "Sheet"
        base = cleaned[:31]
        candidate = base
        suffix = 2
        while candidate.casefold() in used_names:
            marker = f" ({suffix})"
            candidate = f"{base[:31 - len(marker)]}{marker}"
            suffix += 1
        used_names.add(candidate.casefold())
        return candidate

    def _populate_sheet(
        self,
        worksheet,
        items: list[CustomerDebtExportItem],
        filters: CustomerDebtExportFilters,
    ) -> None:
        dark_blue = "17365D"
        header_blue = "2F75B5"
        pale_blue = "D9EAF7"
        light_border = Side(style="thin", color="D9E2F3")

        worksheet.sheet_view.showGridLines = False
        worksheet.freeze_panes = "A5"
        worksheet.sheet_properties.tabColor = header_blue
        worksheet.merge_cells("A1:F1")
        worksheet["A1"] = "DANH SÁCH KHÁCH HÀNG VÀ CÔNG NỢ"
        worksheet["A1"].font = Font(name="Aptos Display", size=16, bold=True, color="FFFFFF")
        worksheet["A1"].fill = PatternFill("solid", fgColor=dark_blue)
        worksheet["A1"].alignment = Alignment(horizontal="center", vertical="center")
        worksheet.row_dimensions[1].height = 28

        worksheet.merge_cells("A2:F2")
        worksheet["A2"] = self._period_label(filters.date_from, filters.date_to)
        worksheet["A2"].font = Font(name="Aptos", size=10, italic=True, color="44546A")
        worksheet["A2"].fill = PatternFill("solid", fgColor=pale_blue)
        worksheet["A2"].alignment = Alignment(horizontal="left", vertical="center")
        worksheet.row_dimensions[2].height = 21

        for column, header in enumerate(self.HEADERS, start=1):
            cell = worksheet.cell(row=4, column=column, value=header)
            cell.font = Font(name="Aptos", size=10, bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor=header_blue)
            cell.alignment = Alignment(horizontal="center", vertical="center")
        worksheet.row_dimensions[4].height = 23

        first_data_row = 5
        for index, item in enumerate(items, start=1):
            row = first_data_row + index - 1
            values = (
                index,
                self._safe_text(item.label),
                self._safe_text(item.ward_name),
                self._safe_text(item.village_name),
                self._safe_text(item.address_detail),
                self._rounded_debt(item.current_debt),
            )
            for column, value in enumerate(values, start=1):
                cell = worksheet.cell(row=row, column=column, value=value)
                cell.font = Font(name="Aptos", size=10)
                cell.border = Border(bottom=light_border)
                cell.alignment = Alignment(
                    horizontal="right" if column in (1, 6) else "left",
                    vertical="center",
                )
            worksheet.cell(row=row, column=6).number_format = '#,##0;[Red]-#,##0'

        last_data_row = first_data_row + len(items) - 1
        total_row = max(first_data_row, last_data_row + 1)
        worksheet.merge_cells(start_row=total_row, start_column=1, end_row=total_row, end_column=5)
        worksheet.cell(row=total_row, column=1, value=f"Tổng cộng ({len(items)} khách hàng)")
        worksheet.cell(row=total_row, column=1).font = Font(name="Aptos", size=10, bold=True, color=dark_blue)
        worksheet.cell(row=total_row, column=1).alignment = Alignment(horizontal="right")
        total_cell = worksheet.cell(row=total_row, column=6)
        total_cell.value = f"=SUM(F{first_data_row}:F{last_data_row})" if items else 0
        total_cell.font = Font(name="Aptos", size=10, bold=True, color=dark_blue)
        total_cell.fill = PatternFill("solid", fgColor=pale_blue)
        total_cell.number_format = '#,##0;[Red]-#,##0'
        total_cell.alignment = Alignment(horizontal="right")

        if items:
            worksheet.auto_filter.ref = f"A4:F{last_data_row}"
        widths = (8, 32, 22, 24, 42, 20)
        for column, width in enumerate(widths, start=1):
            worksheet.column_dimensions[get_column_letter(column)].width = width
        worksheet.auto_filter.ref = f"A4:F{max(4, last_data_row)}"

    @staticmethod
    def _period_label(date_from: date | None, date_to: date | None) -> str:
        if not date_from and not date_to:
            return "Kỳ công nợ: Tất cả thời gian"
        start = date_from.strftime("%d/%m/%Y") if date_from else "đầu kỳ dữ liệu"
        end = date_to.strftime("%d/%m/%Y") if date_to else "hiện tại"
        return f"Kỳ công nợ: {start} - {end} (bao gồm hai đầu ngày)"

    @staticmethod
    def _safe_text(value: str | None) -> str:
        text = value or ""
        return f"'{text}" if text.startswith(("=", "+", "-", "@")) else text

    @staticmethod
    def _rounded_debt(value: Decimal) -> int:
        return int(value.quantize(Decimal("1"), rounding=ROUND_HALF_UP))
