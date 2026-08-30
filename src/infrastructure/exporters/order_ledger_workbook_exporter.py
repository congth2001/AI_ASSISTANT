"""OpenPyXL adapter for the single-customer transaction ledger workbook."""

import re
from datetime import datetime
from decimal import Decimal, ROUND_HALF_UP
from io import BytesIO
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

from src.domain.entities.dashboard_analytics import (
    OrderLedgerExportData,
    OrderLedgerExportFilters,
)
from src.domain.services.i_order_ledger_workbook_exporter import (
    IOrderLedgerWorkbookExporter,
)


class OpenpyxlOrderLedgerWorkbookExporter(IOrderLedgerWorkbookExporter):
    MAX_DATA_ROWS = 1_048_570
    HEADERS = (
        "Ngày",
        "Số phiếu",
        "Lý do",
        "Tên hàng",
        "ĐVT",
        "Đơn giá",
        "Số lượng",
        "Thanh toán",
        "Ghi nợ",
        "Nợ cộng dồn",
    )
    ACCOUNTING_FORMAT = '_(* #,##0_);_(* \\(#,##0\\);_(* "-"??_);_(@_)'
    DATE_TIME_FORMAT = "d/m/yyyy h:mm AM/PM"

    def build(
        self,
        data: OrderLedgerExportData,
        _filters: OrderLedgerExportFilters,
    ) -> bytes:
        if len(data.items) > self.MAX_DATA_ROWS:
            raise ValueError("The order ledger exceeds the Excel row limit")

        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = self._sheet_name(data.customer_name)

        metadata = (
            ("Tên đối tác:", data.customer_name),
            ("Địa chỉ:", data.customer_address or ""),
            (
                "Ngày xuất file:",
                datetime.now(ZoneInfo("Asia/Ho_Chi_Minh")).replace(tzinfo=None),
            ),
        )
        for row_number, (label, value) in enumerate(metadata, start=1):
            label_cell = worksheet.cell(row=row_number, column=1, value=label)
            label_cell.font = Font(name="Arial", size=11, bold=True)
            value_cell = worksheet.cell(
                row=row_number,
                column=2,
                value=self._safe_text(value) if isinstance(value, str) else value,
            )
            value_cell.font = Font(name="Arial", size=11)
            value_cell.alignment = Alignment(horizontal="right")
        worksheet["B3"].number_format = self.DATE_TIME_FORMAT

        for column, header in enumerate(self.HEADERS, start=1):
            cell = worksheet.cell(row=6, column=column, value=header)
            cell.font = Font(name="Arial", size=11, bold=True)

        previous_event_key: tuple | None = None
        row_number = 7
        for item in data.items:
            event_key = (item.occurred_at, item.voucher_number, item.reason)
            first_event_row = event_key != previous_event_key
            previous_event_key = event_key
            occurred_at = item.occurred_at
            if occurred_at.tzinfo is not None:
                occurred_at = occurred_at.replace(tzinfo=None)
            values = (
                occurred_at if first_event_row else None,
                self._safe_text(item.voucher_number) if first_event_row else None,
                item.reason if first_event_row else None,
                self._safe_text(item.product_name),
                self._safe_text(item.unit_name),
                self._rounded_money(item.unit_price),
                item.quantity,
                self._rounded_money(item.payment_amount) if first_event_row else None,
                self._rounded_money(item.debt_amount) if first_event_row else None,
                self._rounded_money(item.running_debt) if first_event_row else None,
            )
            for column, value in enumerate(values, start=1):
                cell = worksheet.cell(row=row_number, column=column, value=value)
                cell.font = Font(name="Arial", size=11)
            worksheet.cell(row=row_number, column=1).number_format = self.DATE_TIME_FORMAT
            for column in range(6, 11):
                worksheet.cell(row=row_number, column=column).number_format = self.ACCOUNTING_FORMAT
            row_number += 1

        if not data.items:
            worksheet.cell(row=7, column=1, value="Không có giao dịch phù hợp với bộ lọc.")
            worksheet["A7"].font = Font(name="Arial", size=11, italic=True)

        widths = (18.875, 25, 8.75, 37, 13, 13.25, 9.25, 14.25, 15, 15.25)
        for column, width in enumerate(widths, start=1):
            worksheet.column_dimensions[chr(64 + column)].width = width

        output = BytesIO()
        workbook.save(output)
        return output.getvalue()

    @staticmethod
    def _sheet_name(customer_name: str) -> str:
        safe_name = re.sub(r"[\\/*?:\[\]]", " ", customer_name).strip() or "DoiTac"
        suffix = "_DoiTac"
        return f"{safe_name[:31 - len(suffix)]}{suffix}"

    @staticmethod
    def _safe_text(value: str | None) -> str:
        text = value or ""
        return f"'{text}" if text.startswith(("=", "+", "-", "@")) else text

    @staticmethod
    def _rounded_money(value: Decimal) -> int:
        return int(value.quantize(Decimal("1"), rounding=ROUND_HALF_UP))
